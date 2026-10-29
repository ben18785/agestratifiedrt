"""Turn the raw downloads in data/raw/ into small CSVs in data/processed/ (committed).

Outputs
  england_cases_by_age.csv      daily cases by specimen date, England, 5-year age bands
  england_cases_total.csv       daily cases by specimen date, England, all ages (for checking)
  comix_participants_by_week.csv  CoMix participants per survey week and age band
  comix_contacts_by_week.csv      CoMix contacts per survey week, participant age band,
                                  contact age band and setting
  ons_cis_daily_by_age.csv      ONS infection survey, modelled daily positivity by age/school year,
                                28 Feb to 24 Apr 2021 (releases of 16 and 30 April 2021)
  ons_cis_14day_by_age.csv      ONS infection survey, non-overlapping 14-day positivity by
                                age/school year, May 2020 to Dec 2021
Usage: python data/prepare_data.py   (after python data/download_data.py)
"""
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
RAW, OUT = HERE / "raw", HERE / "processed"
OUT.mkdir(exist_ok=True)

# ---- UKHSA cases by age ------------------------------------------------------------------
# The archive file holds many snapshots (data releases) of each day's count with no release
# date. As in the ukc19 package, take the largest value for each date and age band: the age
# bands then sum to the final all-ages series (median ratio 0.9997).
c = pd.read_csv(RAW / "ukhsa_cases/nation_newCasesBySpecimenDateAgeDemographics_areas_letter_E.csv")
bands = sorted(a for a in c.age.unique() if a not in ("00_59", "60+", "unassigned"))
cases = (c[c.age.isin(bands)].groupby(["date", "age"], as_index=False).value.max()
         .rename(columns={"age": "age_band", "value": "cases"}))
cases.to_csv(OUT / "england_cases_by_age.csv", index=False)
t = pd.read_csv(RAW / "ukhsa_cases/nation_newCasesBySpecimenDate.csv")
t = t[t.area_name == "England"].groupby("date", as_index=False).value.max().rename(columns={"value": "cases"})
t.to_csv(OUT / "england_cases_total.csv", index=False)
chk = cases.groupby("date").cases.sum().to_frame("age_sum").join(t.set_index("date"))
print("cases: age-band sum / all-ages total, median %.4f" % (chk.age_sum / chk.cases).median())

# ---- CoMix --------------------------------------------------------------------------------
# Unweighted. Contacts reported individually only (contact_common); large group contacts in
# contact_extra are not included. Contact age = midpoint of the reported range. Setting
# priority when several apply: home > school > work > other.
AGE_EDGES = [0, 5, 12, 18, 30, 40, 50, 60, 70, np.inf]
AGE_LABELS = ["0-4", "5-11", "12-17", "18-29", "30-39", "40-49", "50-59", "60-69", "70+"]
part_fix = {"70-120": "70+", "18-19": "18-29", "Under 1": "0-4"}
p = pd.read_csv(RAW / "comix/CoMix_uk_participant_common.csv")
p["part_age"] = p.part_age.replace(part_fix)
p = p[p.part_age.isin(AGE_LABELS)]
s = pd.read_csv(RAW / "comix/CoMix_uk_sday.csv")
s["survey_date"] = pd.to_datetime(dict(year=2000 + s.year, month=s.month, day=s.day))
s["week_start"] = (s.survey_date - pd.to_timedelta(s.survey_date.dt.day_of_week, unit="D")).dt.date
s["panel"] = s.part_id.str[0]
s = s[["part_id", "survey_date", "week_start", "panel", "wave"]].merge(p[["part_id", "part_age"]], on="part_id")
parts = (s.groupby(["week_start", "part_age"], as_index=False)
         .agg(n_participants=("part_id", "size"), panels=("panel", lambda x: "".join(sorted(set(x)))),
              waves=("wave", lambda x: ",".join(map(str, sorted(set(x)))))))
parts.to_csv(OUT / "comix_participants_by_week.csv", index=False)

k = pd.read_csv(RAW / "comix/CoMix_uk_contact_common.csv",
                usecols=["part_id", "cnt_age_exact", "cnt_age_est_min", "cnt_age_est_max",
                         "cnt_home", "cnt_school", "cnt_work"])
age = k.cnt_age_exact.where(k.cnt_age_exact.notna(), (k.cnt_age_est_min + k.cnt_age_est_max) / 2)
k["cnt_age"] = pd.cut(age, AGE_EDGES, right=False, labels=AGE_LABELS)
k["setting"] = np.select([k.cnt_home == True, k.cnt_school == True, k.cnt_work == True],
                         ["home", "school", "work"], "other")
k = k.dropna(subset=["cnt_age"]).merge(s[["part_id", "week_start", "part_age"]], on="part_id")
cont = (k.groupby(["week_start", "part_age", "cnt_age", "setting"], observed=True, as_index=False)
        .size().rename(columns={"size": "n_contacts"}))
cont.to_csv(OUT / "comix_contacts_by_week.csv", index=False)
print("CoMix: %d survey-weeks, %d participants, %d contacts with age" %
      (parts.week_start.nunique(), parts.n_participants.sum(), cont.n_contacts.sum()))


# ---- ONS infection survey -----------------------------------------------------------------
def read_grouped_table(path, sheet, header_row, sub_rows, first_data_row):
    """Tables with one column group per age band: the band names sit on header_row and the
    measure names on the rows below. Returns a long data frame."""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    head = rows[header_row]
    subs = [rows[r] for r in sub_rows]
    group, out = None, []
    cols = []
    for j in range(1, len(head)):
        group = head[j] if head[j] else group
        name = " ".join(str(sr[j]) for sr in subs if j < len(sr) and sr[j]) or ""
        cols.append((j, group, name))
    for r in rows[first_data_row:]:
        if r[0] is None or not any(r[1:]):
            break
        for j, g, name in cols:
            if g is not None and j < len(r):
                out.append(dict(period=r[0], age_group=" ".join(str(g).split()), measure=name, value=r[j]))
    return pd.DataFrame(out)


daily = []
for f in ["ons_cis_20210416.xlsx", "ons_cis_20210430.xlsx"]:
    wb = openpyxl.load_workbook(RAW / "ons_cis" / f, read_only=True)
    sheet = next(ws.title for ws in wb.worksheets
                 if "daily rates" in str(next(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]).lower()
                 and "age/school year" in str(next(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]).lower())
    d = read_grouped_table(RAW / "ons_cis" / f, sheet, 4, [5, 6], 7)
    d["release"] = f[8:16]
    daily.append(d)
daily = pd.concat(daily)
daily["measure"] = daily.measure.map(lambda m: "lower" if "Lower" in m else "upper" if "Upper" in m else "estimate")
daily = daily.pivot_table(index=["release", "period", "age_group"], columns="measure", values="value",
                          aggfunc="first").reset_index()
for col in ["estimate", "lower", "upper"]:
    daily[col] = pd.to_numeric(daily[col], errors="coerce")
    if daily[col].max() < 1:          # these releases give proportions, despite the "%" header
        daily[col] *= 100
daily = daily.rename(columns={"period": "date", "estimate": "pct_positive", "lower": "pct_lower", "upper": "pct_upper"})
daily["date"] = pd.to_datetime(daily.date).dt.date
daily.to_csv(OUT / "ons_cis_daily_by_age.csv", index=False)

f14 = read_grouped_table(RAW / "ons_cis/ons_cis_2021_edition.xlsx", "1j", 4, [5], 6)
f14["measure"] = f14.measure.map(lambda m: "pct_lower" if "Lower" in m else "pct_upper" if "Upper" in m
                                 else "n_positive" if "Number of people testing" in m
                                 else "n_sample" if "sample" in m else "pct_positive")
f14 = f14.pivot_table(index=["period", "age_group"], columns="measure", values="value", aggfunc="first").reset_index()
dates = f14.period.str.split(" to ", expand=True)
f14.insert(0, "start", pd.to_datetime(dates[0], dayfirst=True).dt.date)
f14.insert(1, "end", pd.to_datetime(dates[1], dayfirst=True).dt.date)
f14 = f14.drop(columns="period").sort_values(["start", "age_group"])
f14.to_csv(OUT / "ons_cis_14day_by_age.csv", index=False)
print("ONS: %d daily rows (%s to %s), %d fortnightly rows" %
      (len(daily), daily.date.min(), daily.date.max(), len(f14)))
